#! python3
import openpyxl, docx
boss = openpyxl.load_workbook('xuan.xlsx')
sheet = boss.get_sheet_by_name('Sheet2')
sheet.freeze_panes = 'A2'
doc = docx.Document('anqiao6_22.docx')
yewu= 6
tixing =1
start = 2
Text=[]
for para in doc.paragraphs:
    Text.append(para.text)
for i in range(start, len(Text)+start):
    sheet['A' + str(i)] = yewu
    sheet['B' + str(i)] = tixing
    sheet['C' + str(i)] = int(i+1-start)
    sheet['D' + str(i)] = Text[i-start][len(str(i))+1:-3]
    if Text[i-start][-2] == '×':
        sheet['E' + str(i)] = 2
    else:
        sheet['E' + str(i)] = 1
#doc.save('anqiao1.docx')
boss.save('xuan.xlsx')

