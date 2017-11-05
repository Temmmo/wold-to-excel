# -*- coding: utf-8 -*-
"""
Created on Thu Aug 17 11:52:49 2017

@author: ljd19
"""

import docx, openpyxl
#os.getcwd()

doc = docx.Document('anqiao.docx')

table = openpyxl.load_workbook('test2.xlsx')

sheet = table.get_sheet_by_name('Sheet1')

Text=[]
start =2
for para in doc.paragraphs:
    Text.append(para.text)
for i in range(start, len(Text)+start):
    sheet['A' + str(i)] = Text[0]



