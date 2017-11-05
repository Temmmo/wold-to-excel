#! python3
import openpyxl, docx
boss = openpyxl.load_workbook('tuo.xlsx')
sheet = boss.get_sheet_by_name('Sheet2')
sheet.freeze_panes = 'A2'
doc = docx.Document('anqiao6_22.docx')
Text=[]
yewu= 3
tixing =2
start =2
for para in doc.paragraphs:
    Text.append(para.text)
Num = []
for i in range(2, len(Text)):
    if Text[i-1][1] == 'C':
        answer = Text[i-2]+Text[i-1]
        del Text[i-2]
        Num.append(i-1)
        Text.insert(i-2, answer)
for i in range(0,len(Num)):
    del Text[Num[i]]
    for j in range(i+1,len(Num)):
        Num[j] = Num[j]-1
number = int(len(Text)/2)
for i in range(start-1, number+start-1):
    a = Text[i * 2 - 1].find('A')
    b = Text[i * 2 - 1].find('B')
    c = Text[i * 2 - 1].find('C')
    d = Text[i * 2 - 1].find('D')
    e = int(len(Text[i * 2 - 1]))
    for j in range((i-1)*4,i*4):
        sheet['A' + str(j+start)] = yewu
        sheet['B' + str(j+start)] = tixing
        sheet['C' + str(j+start)] = int(i)
        sheet['D' + str(j+start)] = int(j % 4+1)
        if int(j % 4 + 1) == 1:
            sheet['E' + str(j + start)] = 'A'
            sheet['F' + str(j + start)] =Text[i * 2 - 1][a+2:b-1]
        elif int(j % 4 + 1) == 2:
            sheet['E' + str(j+ start)] = 'B'
            sheet['F' + str(j + start)] = Text[i * 2 - 1][b + 2:c-1]
        elif int(j % 4 + 1) == 3:
            sheet['E' + str(j+start)] = 'C'
            sheet['F' + str(j + start)] = Text[i * 2 - 1][c + 2:d-1]
        else:
            sheet['E' + str(j+start)] = 'D'
            sheet['F' + str(j + start)] = Text[i * 2 - 1][d + 2:e]
boss.save('tuo.xlsx')