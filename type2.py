#! python3
import openpyxl, docx, os, pyperclip
import re
boss = openpyxl.load_workbook('xuan.xlsx')
sheet = boss.get_sheet_by_name('Sheet1')
sheet.freeze_panes = 'A2'
doc = docx.Document('anqiao6_22.docx')
Text=[]
yewu= 3
tixing =2
start = 2
for para in doc.paragraphs:
    Text.append(para.text)
Num =[]
for i in range(2, len(Text)):
    if Text[i-1][1] == 'C':
        answer = Text[i-2]+Text[i-1]
        del Text[i-2]
        Num.append(i-1)
        Text.insert(i-2, answer)
for i in range(0 ,len(Num)):
    del Text[Num[i]]
    for j in range(i+1,len(Num)):
        Num[j] = Num[j]-1
a = int(len(Text)/2)
for i in range(start, a+start):
    sheet['A' + str(i)] = yewu
    sheet['B' + str(i)] = tixing
    sheet['C' + str(i)] = int(i-start+1)
    options = re.compile(r'[ABCD]')
    an = options.search(Text[(i-start)*2])
    if an.group(0) == 'A':
        sheet['E' + str(i)] = 'A'
    elif an.group(0) == 'B':
        sheet['E' + str(i)] = 'B'
    elif an.group(0) == 'C':
        sheet['E' + str(i)] = 'C'
    else:
        sheet['E' + str(i)] = 'D'
    m = Text[(i-start)*2].find(an.group(0))
    n = int(len(Text[(i-start)*2]))
    sheet['D' + str(i)] = Text[(i-start)*2][len(str(i)) + 1:m]+Text[(i-start)*2][m+1:n]
boss.save('xuan.xlsx')